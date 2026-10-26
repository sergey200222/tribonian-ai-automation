"""Offline unit tests. No external API or production data."""
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import fitz

SCRIPT = Path(__file__).resolve().parents[1] / 'summarize_folder.py'
spec = importlib.util.spec_from_file_location('summarize_folder', SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def test_chunk_text_preserves_text():
    text = 'abc\n\n' * 9000
    chunks = module.chunk_text(text, 1000)
    assert len(chunks) > 1
    assert ''.join(chunks).replace('\n', '') == text.replace('\n', '')
    assert all(len(c) <= 1002 for c in chunks)


def test_supported_files_and_recursion(tmp_path):
    (tmp_path / 'folder').mkdir()
    (tmp_path / 'folder' / 'one.PDF').write_bytes(b'not a PDF')
    (tmp_path / 'two.txt').write_text('x', encoding='utf-8')
    (tmp_path / 'bad.exe').write_text('bad', encoding='utf-8')
    (tmp_path / '.hidden').mkdir()
    (tmp_path / '.hidden' / 'secret.txt').write_text('dont scan')
    supported, skipped = module.find_files(tmp_path)
    assert [module.path_key(x, tmp_path) for x in supported] == ['folder/one.PDF', 'two.txt']
    assert [x.name for x in skipped] == ['bad.exe']


def test_pdf_extracts_text_and_keeps_page_numbers(tmp_path):
    path = tmp_path / 'fixture.pdf'
    doc = fitz.open()
    for page_number in (1, 2):
        page = doc.new_page()
        page.insert_text((40, 40), f'Article {page_number} ' + 'Example clause ' * 40)
    doc.save(path)
    doc.close()
    extracted, record = module.extract_file(path, tmp_path, None, 'gpt-4.1-mini', check=True)
    assert record['pages'] == 2
    assert record['status'] == 'ok'
    assert 'страница PDF 1' in extracted and 'страница PDF 2' in extracted
    assert 'Article 1' in extracted and 'Article 2' in extracted


def test_image_check_does_not_call_api(tmp_path):
    path = tmp_path / 'sample.webp'
    path.write_bytes(b'synthetic')
    extracted, record = module.extract_file(path, tmp_path, None, 'gpt-4.1-mini', check=True)
    assert record['status'] == 'needs_vision'
    assert 'Vision' in extracted


def test_plain_text_and_docx(tmp_path):
    from docx import Document
    text_file = tmp_path / 'content.txt'
    text_file.write_text('\ufeffРусский текст', encoding='utf-8')
    text, record = module.extract_file(text_file, tmp_path, None, 'fake', check=True)
    assert text == 'Русский текст' and record['status'] == 'ok'
    path = tmp_path / 'sample.docx'
    doc = Document()
    doc.add_paragraph('Глава 1')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Статья'
    table.cell(0, 1).text = 'Содержание'
    doc.save(path)
    text, record = module.extract_file(path, tmp_path, None, 'fake', check=True)
    assert 'Глава 1' in text
    assert 'Статья | Содержание' in text


def test_ask_text_uses_no_response_storage():
    class FakeResponses:
        def create(self, **kw):
            assert kw['store'] is False
            assert kw['model'] == 'unit-test-model'
            return SimpleNamespace(output_text='Test summary')
    client = SimpleNamespace(responses=FakeResponses())
    assert module.ask_text(client, 'unit-test-model', 'instructions', 'content') == 'Test summary'


def test_dry_run_outputs_manifest_without_api_key(tmp_path, monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    inp = tmp_path / 'source'
    inp.mkdir()
    (inp / 'page.jpg').write_bytes(b'mock image')
    (inp / 'notes.txt').write_text('Notes', encoding='utf-8')
    dest = tmp_path / 'result'
    code = module.main(['--input', str(inp), '--output', str(dest), '--check'])
    assert code == 0
    manifest = json.loads((dest / 'manifest.json').read_text(encoding='utf-8'))
    assert manifest['mode'] == 'check'
    assert len(manifest['files']) == 2
    assert {row['status'] for row in manifest['files']} == {'ok', 'needs_vision'}
    assert not (dest / 'summary.md').exists()


def test_unsupported_file_return_code(tmp_path):
    inp = tmp_path / 'in'
    inp.mkdir()
    (inp / 'ok.txt').write_text('Here', encoding='utf-8')
    (inp / 'not-supported.bin').write_text('data', encoding='utf-8')
    out = tmp_path / 'out'
    assert module.main(['--input', str(inp), '--output', str(out), '--check']) == 1
    manifest = json.loads((out / 'manifest.json').read_text(encoding='utf-8'))
    assert any(r['status'] == 'unsupported' for r in manifest['files'])
